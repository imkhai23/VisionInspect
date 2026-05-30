import sys
import subprocess
import os

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(13)

for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = True
    return h

# Cover Page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("TRƯỜNG ĐẠI HỌC ĐÀ LẠT\nKHOA CÔNG NGHỆ THÔNG TIN")
run.bold = True
run.font.size = Pt(14)

doc.add_paragraph("\n\n\n\n")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("TIỂU LUẬN")
run.bold = True
run.font.size = Pt(24)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Môn học: Các Công Nghệ Mới Trong Phát Triển Phần Mềm")
run.font.size = Pt(14)

doc.add_paragraph("\n\n")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("ĐỀ TÀI:\nHỆ THỐNG KIỂM TRA CHẤT LƯỢNG SẢN PHẨM TRONG NHÀ MÁY BẰNG TRÍ TUỆ NHÂN TẠO")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph("\n\n\n\n")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p.add_run("Giáo viên hướng dẫn: Nguyễn Trọng Hiếu\n").bold = True
p.add_run("Sinh viên thực hiện: Phan Thanh Khải\n").bold = True
p.add_run("MSSV: 2212386").bold = True

doc.add_paragraph("\n\n\n\n\n\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Đà Lạt, 2026").italic = True

doc.add_page_break()

# Content
add_heading(doc, 'TÓM TẮT NỘI DUNG ĐỀ TÀI', 1)

add_heading(doc, '1. Mục tiêu chính của đề tài', 2)
doc.add_paragraph("- Xây dựng pipeline AI thời gian thực sử dụng YOLOv8 để phát hiện và phân loại lỗi sản phẩm.")
doc.add_paragraph("- Tích hợp thuật toán ByteTrack để theo dõi đối tượng qua các frame, đếm số lượng sản phẩm không bị trùng lặp.")
doc.add_paragraph("- Truyền video trực tiếp qua giao thức MJPEG và gửi siêu dữ liệu (metadata) qua WebSocket để giám sát thời gian thực trên giao diện web.")
doc.add_paragraph("- Xây dựng hệ thống tự động ghi hình (Smart Recording) và gửi cảnh báo qua Telegram Bot khi phát hiện lỗi.")
doc.add_paragraph("- Phát triển nền tảng quản trị (Admin Platform) để quản lý bộ dữ liệu (dataset) và trực tiếp huấn luyện lại mô hình AI hoàn chỉnh trên nền web.")
doc.add_paragraph("- Triển khai toàn bộ hệ thống bằng Docker Compose với 4 services độc lập, sẵn sàng hoạt động dưới dạng phần mềm dịch vụ SaaS.")

add_heading(doc, '2. Bảng mô tả ngắn các công nghệ sử dụng', 2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Thành phần'
hdr_cells[1].text = 'Công nghệ'
hdr_cells[2].text = 'Mục đích'
techs = [
    ('Frontend', 'Next.js 14, TypeScript', 'Xây dựng giao diện người dùng, xử lý App Router, SSR'),
    ('Backend', 'FastAPI, Python', 'Xây dựng REST API, WebSocket, AI Inference bất đồng bộ'),
    ('AI Engine', 'YOLOv8, ByteTrack', 'Phát hiện lỗi và theo dõi đối tượng đa mục tiêu'),
    ('Cơ sở dữ liệu', 'Supabase (PostgreSQL)', 'Xác thực (Auth), lưu trữ dữ liệu (Database) và tệp tin (Storage)'),
    ('Giao tiếp Realtime', 'WebSocket, MJPEG', 'Streaming video liên tục và gửi metadata với độ trễ thấp'),
    ('Hàng đợi (Queue)', 'Redis, RQ', 'Xử lý bất đồng bộ (background jobs) cho các tác vụ huấn luyện AI'),
    ('Triển khai', 'Docker, Docker Compose', 'Đóng gói và triển khai toàn bộ ứng dụng một cách độc lập')
]
for t1, t2, t3 in techs:
    row_cells = table.add_row().cells
    row_cells[0].text = t1
    row_cells[1].text = t2
    row_cells[2].text = t3

add_heading(doc, '3. Công cụ AI Agents đã sử dụng', 2)
doc.add_paragraph("Trong quá trình phát triển dự án, các công cụ Trí tuệ nhân tạo (LLMs) đã được ứng dụng sâu rộng để tăng tốc độ và cải thiện chất lượng phần mềm:")
doc.add_paragraph("- GitHub Copilot: Đóng vai trò là người trợ lý lập trình. Cung cấp tính năng gợi ý mã nguồn (Code Completion) tự động hoàn thành logic, sinh code tự động dựa trên comment định hướng (Code Generation), và hỗ trợ tìm/sửa lỗi nhanh chóng (Debugging).")
doc.add_paragraph("- ChatGPT / LLMs: Đóng vai trò là nhà tư vấn giải pháp. Hỗ trợ nghiên cứu, so sánh các công nghệ mới (ví dụ: MJPEG so với WebRTC), thiết kế sơ đồ kiến trúc hệ thống, đồng thời hỗ trợ rà soát văn phong, cấu trúc trong quá trình viết báo cáo, tài liệu kỹ thuật.")

add_heading(doc, '4. Bảng phân tích chức năng/module', 2)
table2 = doc.add_table(rows=1, cols=2)
table2.style = 'Table Grid'
h2 = table2.rows[0].cells
h2[0].text = 'Module / Layer'
h2[1].text = 'Chi tiết chức năng'
modules = [
    ('Client Layer (Frontend)', 'Cung cấp giao diện trực quan cho người dùng. Bao gồm: Dashboard tổng quan, trang giám sát Realtime, lịch sử kiểm tra (History), kiểm tra ảnh tĩnh (Inspect) và khu vực quản trị AI (Admin Platform).'),
    ('API Layer (Backend)', 'Xử lý các logic nghiệp vụ, quản lý middleware xác thực JWT từ Supabase, cung cấp RESTful API và cấu hình WebSocket/MJPEG stream.'),
    ('AI & Video Layer', 'Chạy vòng lặp lấy frame từ camera (VideoProcessor), đẩy vào pipeline xử lý YOLOv8 (VisionEngine) để nhận diện lỗi, crop ảnh lỗi và gọi dịch vụ cảnh báo (AlertService).'),
    ('Data Layer (Supabase)', 'Quản lý toàn vẹn dữ liệu hệ thống thông qua PostgreSQL, bảo mật quyền truy cập với RLS (Row Level Security) và lưu trữ ảnh chụp lỗi/model qua Storage.'),
    ('Queue Layer', 'Hoạt động độc lập với Redis làm message broker. Lắng nghe và chạy ngầm (RQ worker) các tiến trình huấn luyện mô hình tốn nhiều tài nguyên mà không làm chậm máy chủ chính.')
]
for m1, m2 in modules:
    row_cells = table2.add_row().cells
    row_cells[0].text = m1
    row_cells[1].text = m2

add_heading(doc, '5. Bảng phân tích CSDL (Supabase)', 2)
table3 = doc.add_table(rows=1, cols=2)
table3.style = 'Table Grid'
h3 = table3.rows[0].cells
h3[0].text = 'Tên bảng'
h3[1].text = 'Mô tả dữ liệu'
db_tables = [
    ('users', 'Lưu trữ thông tin định danh (id, email, full_name), quyền hạn (is_admin) và trạng thái của người dùng.'),
    ('predictions', 'Lưu lại lịch sử các lần dự đoán lỗi sản phẩm. Gồm user_id, tên ảnh, nhãn lỗi (label), độ tin cậy (confidence) và thời gian xử lý.'),
    ('subscriptions', 'Quản lý gói cước đăng ký của khách hàng (Free/Pro), trạng thái gói và chu kỳ thanh toán nhằm phục vụ mô hình kinh doanh SaaS.'),
    ('usage_logs', 'Ghi nhận chi tiết từng thao tác của người dùng để thống kê mức độ sử dụng hàng tháng.'),
    ('datasets', 'Quản lý metadata của các bộ dữ liệu do Admin tải lên (tên, mô tả, danh sách các nhãn/classes).'),
    ('training_jobs', 'Lưu trữ trạng thái tiến trình (queued, running, completed), thông số huấn luyện, giúp theo dõi Job qua giao diện web.'),
    ('model_versions', 'Lưu trữ lịch sử các phiên bản mô hình AI đã hoàn thành, vị trí lưu file (.pt) và chỉ số mAP, cho phép kích hoạt/rollback mô hình dễ dàng.')
]
for d1, d2 in db_tables:
    row_cells = table3.add_row().cells
    row_cells[0].text = d1
    row_cells[1].text = d2

add_heading(doc, '6. Bảng mô tả các chức năng đã hoàn thiện', 2)
table4 = doc.add_table(rows=1, cols=2)
table4.style = 'Table Grid'
h4 = table4.rows[0].cells
h4[0].text = 'Chức năng'
h4[1].text = 'Mô tả chi tiết'
features = [
    ('Giám sát thời gian thực', 'Hiển thị luồng video camera trực tiếp, vẽ bounding box bao quanh sản phẩm, đếm số lượng tốt/lỗi và cập nhật thống kê mà không cần làm mới trang.'),
    ('Kiểm tra thủ công', 'Cho phép người dùng upload một bức ảnh sản phẩm tĩnh để AI phân tích và trả về kết quả ngay lập tức.'),
    ('Lịch sử & Thống kê', 'Liệt kê danh sách các ảnh đã kiểm tra, hiển thị biểu đồ trực quan phân bố các loại lỗi và hiển thị hạn mức sử dụng gói cước của người dùng.'),
    ('Quản trị Dataset', 'Tạo bộ dữ liệu mới, upload file ảnh và nhãn dạng ZIP. Hệ thống tự động giải nén và cấu trúc hóa để chuẩn bị huấn luyện.'),
    ('Huấn luyện Model', 'Cấu hình số epochs, chọn dataset và khởi chạy Job huấn luyện YOLOv8. Tiến độ cập nhật realtime lên giao diện Admin qua WebSocket.'),
    ('Deploy Model nóng', 'Theo dõi danh sách các mô hình hiện có, cho phép Admin nhấn nút "Deploy" để chuyển đổi mô hình đang chạy mà không làm gián đoạn Backend.'),
    ('Cảnh báo Telegram', 'Hệ thống tự động cắt khung hình chứa sản phẩm lỗi, ghép chung với thông tin độ tin cậy và thời gian để gửi cảnh báo trực tiếp về ứng dụng Telegram.')
]
for f1, f2 in features:
    row_cells = table4.add_row().cells
    row_cells[0].text = f1
    row_cells[1].text = f2

add_heading(doc, '7. Kết luận, tự nhận xét, đánh giá', 2)
p_concl_1 = doc.add_paragraph()
p_concl_1.add_run("Kết luận: ").bold = True
p_concl_1.add_run("Đề tài đã hoàn thành xuất sắc mục tiêu xây dựng một nền tảng SaaS hoàn chỉnh ứng dụng AI vào kiểm tra chất lượng sản phẩm. Việc kết hợp thành công các công nghệ hiện đại như FastAPI, Next.js, YOLOv8 và Supabase mang lại một giải pháp có tính ứng dụng thực tiễn cao, đáp ứng chặt chẽ yêu cầu về tốc độ xử lý thời gian thực.")

p_concl_2 = doc.add_paragraph()
p_concl_2.add_run("Tự nhận xét & Đánh giá:\n").bold = True
p_concl_2.add_run("- Điểm mạnh: ").bold = True
p_concl_2.add_run("Kiến trúc phần mềm được thiết kế theo hướng microservices-like linh hoạt, tách biệt luồng I/O camera và tính toán AI (I/O Bound vs CPU/GPU Bound). Việc sử dụng Docker giúp triển khai hệ thống dễ dàng trên mọi môi trường. Đặc biệt, hệ thống tự động hóa hoàn toàn quy trình MLOps (quản lý dataset, huấn luyện và deploy model) trực tiếp trên nền web.\n")
p_concl_2.add_run("- Hạn chế: ").bold = True
p_concl_2.add_run("Hệ thống hiện tại mới chỉ hỗ trợ theo dõi một luồng camera duy nhất (single-camera). Tốc độ xử lý FPS phụ thuộc khá lớn vào phần cứng (cần cấu hình có GPU) để đạt hiệu năng tối đa. Các model weight hiện tại chưa được mã hóa kỹ càng.\n")
p_concl_2.add_run("- Hướng phát triển: ").bold = True
p_concl_2.add_run("Trong tương lai, giải pháp có thể được tối ưu để chạy cục bộ trên các thiết bị Edge Computing (như Jetson Nano hoặc Raspberry Pi), hỗ trợ cấu trúc nhiều camera (Multi-Camera) đồng thời và tích hợp trực tiếp với hệ thống PLC/SCADA để phát tín hiệu điều khiển cánh tay robot tự động loại bỏ sản phẩm lỗi trên băng chuyền.")

output_path = r'd:\VisionInspect\Tieu_Luan_VisionInspect.docx'
doc.save(output_path)
print(f"Document generated successfully at {output_path}")
